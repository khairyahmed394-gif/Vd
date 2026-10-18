import asyncio, base64, assets as A
from playwright.async_api import async_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
NAVY,TEAL="#101C2C","#19C6B4"
logo=base64.b64encode(open("../img/logo_full.png","rb").read()).decode()
def page(ar,title,sub,body):
    d="rtl" if ar else "ltr"; f="IBM Plex Sans Arabic" if ar else "Manrope"
    return f"""<html dir="{d}"><meta charset=utf-8><style>
@page{{size:A4;margin:0}}*{{box-sizing:border-box}}body{{margin:0;font:13px/1.55 '{f}',sans-serif;color:{NAVY}}}
.hd{{background:{NAVY};color:#fff;padding:34px 44px}}.hd img{{height:30px}}h1{{margin:18px 0 4px;font-size:28px}}
.hd p{{margin:0;color:{TEAL};font-size:14px}}.b{{padding:26px 44px}}
.q{{display:flex;gap:12px;align-items:center;border-bottom:1px solid #DDE2E8;padding:9px 0}}
.q i{{flex:none;width:26px;height:26px;border-radius:50%;background:{TEAL};color:{NAVY};font-style:normal;font-weight:700;text-align:center;line-height:26px}}
.q span{{flex:1}}.q em{{flex:none;color:#536273;font-style:normal;font-size:12px;letter-spacing:2px}}
h2{{font-size:15px;color:{TEAL};margin:18px 0 6px}}.box{{background:#F5F7FA;border-radius:10px;padding:12px 16px;margin:8px 0}}
.box.t{{border-inline-start:4px solid {TEAL}}}.ft{{padding:0 44px;color:#536273;font-size:11px}}
table{{width:100%;border-collapse:collapse}}td,th{{border:1px solid #DDE2E8;padding:8px;text-align:start}}th{{background:{NAVY};color:#fff}}
</style><div class=hd><img src="data:image/png;base64,{logo}"><h1>{title}</h1><p>{sub}</p></div><div class=b>{body}</div>
<div class=ft>NEXORA · {"ابنِ أفضل. انمُ بذكاء." if ar else "Build better. Grow smarter."}</div></html>"""
def score(ar):
    i=1 if ar else 0
    q="".join(f'<div class=q><i>{n}</i><span>{x[2] if ar else x[1]}</span><em>{"٠ ١ ٢" if ar else "0 · 1 · 2"}</em></div>' for n,x in enumerate(A.SCORE_Q,1))
    k=("٠ = لا · ١ = جزئيًا · ٢ = نعم" if ar else "0 = No · 1 = Partly · 2 = Yes")
    b="".join(f'<div class="box t"><b>{x[0]} — {x[2] if ar else x[1]}</b><br>{x[3] if not ar else ""}</div>' for x in A.SCORE_BANDS)
    if ar:
        b="".join(f'<div class="box t"><b>{x[0]} — {x[2]}</b></div>' for x in A.SCORE_BANDS)
    return page(ar,"مقياس صحة الأعمال" if ar else "Business Health Scorecard","10 أسئلة · 5 دقائق" if ar else "10 questions · 5 minutes",
      f'<p>{k}</p>{q}<h2>{"نتيجتك (من ٢٠)" if ar else "Your score (out of 20)"}</h2>{b}')
def sop(ar):
    steps=["استقبل الاستفسار خلال ساعتين","سجّل الاسم والمصدر والاحتياج","أرسل رد قياسي مع العرض","حدّد موعد متابعة خلال ٤٨ ساعة","سجّل النتيجة: تم/لم يتم"] if ar else ["Reply to the enquiry within 2 hours","Log name, source and need","Send the standard reply with the offer","Set a follow-up within 48 hours","Record the outcome: won / lost"]
    rows="".join(f"<tr><td>{n}</td><td>{s}</td><td>[ ]</td></tr>" for n,s in enumerate(steps,1))
    H=("الخطوة","الوصف","المسؤول") if ar else ("Step","Action","Owner")
    f=("الغرض","المسؤول","المحفّز","المخرج","المراجعة") if ar else ("Purpose","Owner","Trigger","Output","Review date")
    fl="".join(f'<div class=box><b>{x}:</b> [ ]</div>' for x in f)
    return page(ar,"قالب إجراءات العمل (SOP)" if ar else "SOP Starter Template","مثال: الرد على استفسار جديد" if ar else "Example: Answering a new enquiry",
      f"<h2>{'البيانات' if ar else 'Details'}</h2>{fl}<h2>{'الخطوات' if ar else 'Steps'}</h2><table><tr><th>#</th><th>{H[1]}</th><th>{H[2]}</th></tr>{rows}</table>")
def check(ar):
    it=["رد خلال ساعتين","سجّل كل عميل محتمل","ابعت عرضًا واضحًا بسعر","تابع بعد ٤٨ ساعة","تابع مرة ثانية بعد أسبوع","اسأل عن سبب الرفض","حدّث الأرقام الأسبوعية"] if ar else ["Reply within 2 hours","Log every lead","Send a clear priced offer","Follow up after 48 hours","Follow up again after 1 week","Ask the reason for a 'no'","Update weekly numbers"]
    return page(ar,"قائمة متابعة المبيعات" if ar else "Sales Follow-up Checklist","٧ خطوات" if ar else "7 steps",
      "".join(f"<div class=q><i>☐</i><span>{x}</span></div>" for x in it))
def xlsx():
    wb=openpyxl.Workbook(); s=wb.active; s.title="Scorecard"
    s.append(["#","Area","Question EN","السؤال","Score 0/1/2"])
    for n,x in enumerate(A.SCORE_Q,1): s.append([n,x[0],x[1],x[2],None])
    s["A13"]="Total"; s["E13"]="=SUM(E2:E11)"
    s["A14"]="Band"; s["E14"]='=IF(E13<=7,"Foundations needed",IF(E13<=14,"Growing, uneven","Strong base"))'
    k=wb.create_sheet("Weekly KPI"); k.append(["Metric","Last week","This week","Trend","Owner","Next action"])
    for n,m in enumerate(["New enquiries","Conversion rate","Repeat / retention","Average value per customer","Margin"],2):
        k.append([m,None,None,f'=IF(OR(B{n}="",C{n}=""),"",IF(C{n}>B{n},"▲",IF(C{n}<B{n},"▼","■")))',None,None])
    for w in (s,k):
        for c in w[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="101C2C")
        for col in "ABCDEF": w.column_dimensions[col].width=28
    s.column_dimensions["C"].width=60; s.column_dimensions["D"].width=45
    wb.save("out/magnets/NEXORA_Scorecard_and_KPI_Dashboard.xlsx")
async def main():
    async with async_playwright() as p:
        b=await p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome"); pg=await b.new_page()
        for n,fn in (("Business_Health_Scorecard",score),("SOP_Starter_Template",sop),("Sales_Followup_Checklist",check)):
            for ar in (0,1):
                await pg.set_content(fn(ar)); await pg.wait_for_timeout(400)
                await pg.pdf(path=f"out/magnets/NEXORA_{n}_{'AR' if ar else 'EN'}.pdf",format="A4",print_background=True)
        await b.close()
xlsx(); asyncio.run(main())
